"""Word (.docx) 解析器。

两处关键设计：

  1. **表格转「表头: 值」**：需求文档里的字段约束（范围/长度/枚举）绝大多数写在
     表格中，保留表格结构能显著提升后续字段约束抽取的准确率；
  2. **内嵌图提取（V2.0）**：正文里的原型图、流程图常带关键信息。这里把图片抠成
     资产，并在**原位置**留下 ``[图片 N]`` 占位标记 —— 识别结果之后会回填到该位置，
     图片与上下文文字的相对顺序不会丢。图片本身不做识别，因此本模块不依赖任何
     OCR 引擎。

遍历按 body 的**子元素顺序**进行（段落与表格交替），而不是先段落后表格：
否则图片占位会全部堆到文末，失去「这张图挨着哪段需求」这条最有价值的线索。
"""
from __future__ import annotations

import io

from app.core.exceptions import DocumentParseError
from app.core.utils import truncate
from app.parsers.base import ASSET_DOCX, DocumentParser, ParsedAsset, ParsedDocument
from app.vision.imaging import extension_for, read_size

try:  # pragma: no cover - 取决于运行环境是否安装 python-docx
    from docx import Document as _DocxDocument

    DOCX_AVAILABLE = True
except ImportError:  # pragma: no cover
    _DocxDocument = None  # type: ignore[assignment]
    DOCX_AVAILABLE = False

#: OOXML 关系命名空间：图片引用（r:embed / r:id / r:link）挂在这里。
RELATIONSHIP_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
RELATIONSHIP_ATTRS = (
    f"{{{RELATIONSHIP_NS}}}embed",
    f"{{{RELATIONSHIP_NS}}}id",
    f"{{{RELATIONSHIP_NS}}}link",
)


class DocxParser(DocumentParser):
    name = "docx"
    label = "Word 解析器"
    extensions = ("docx",)

    def parse(self, data: bytes, *, filename: str = "") -> ParsedDocument:
        if not DOCX_AVAILABLE:  # pragma: no cover
            raise DocumentParseError(
                "解析 .docx 需要安装 python-docx，请执行：pip install python-docx",
                detail={"missing_dependency": "python-docx"},
            )

        try:
            document = _DocxDocument(io.BytesIO(data))
        except Exception as exc:  # noqa: BLE001 - 第三方库异常类型不稳定，统一兜底
            raise DocumentParseError(
                f"Word 文档打开失败，文件可能已损坏或受密码保护：{exc}",
                detail={"filename": filename, "error": str(exc)},
            ) from exc

        lines, table_count, assets, warnings = self._render(document)
        if table_count:
            warnings.insert(
                0,
                f"文档包含 {table_count} 张表格，已按「表头: 值」形式平铺，便于抽取字段约束。",
            )
        if assets:
            warnings.append(
                f"文档包含 {len(assets)} 张内嵌图片，已提取待识别（识别结果会回填到正文的图片位置）。"
            )

        return self._finalize(
            "\n".join(lines),
            extension="docx",
            table_count=table_count,
            metadata={"filename": filename, "paragraph_count": len(document.paragraphs)},
            warnings=warnings,
            assets=assets,
        )

    # ------------------------------------------------------------------ #
    def _render(self, document) -> tuple[list[str], int, list[ParsedAsset], list[str]]:
        """按 body 子元素顺序渲染正文，并就地插入图片占位。"""
        paragraphs = {id(item._p): item for item in document.paragraphs}
        tables = {id(item._tbl): item for item in document.tables}
        related = getattr(document.part, "related_parts", {}) or {}

        lines: list[str] = []
        assets: list[ParsedAsset] = []
        warnings: list[str] = []
        table_count = 0
        last_text = ""

        for element in document.element.body.iterchildren():
            tag = _local_name(element)
            if tag == "p":
                paragraph = paragraphs.get(id(element))
                if paragraph is not None:
                    text = (paragraph.text or "").strip()
                    if text:
                        lines.append(f"{self._heading_prefix(paragraph)}{text}")
                        last_text = text
                anchor = (paragraph.text if paragraph is not None else "") or last_text
            elif tag == "tbl":
                table = tables.get(id(element))
                if table is not None:
                    table_count += 1
                    lines.append("")
                    lines.append(f"[表格 {table_count}]")
                    lines.extend(self._render_table(table))
                anchor = last_text
            else:
                continue

            for rel_id in _image_rel_ids(element):
                asset = _build_asset(len(assets) + 1, rel_id, related, anchor, warnings)
                if asset is None:
                    continue
                assets.append(asset)
                lines.append(asset.placeholder)
        return lines, table_count, assets, warnings

    @staticmethod
    def _heading_prefix(paragraph) -> str:
        """把 Word 标题样式还原成 Markdown 的 #，供规则解析器识别章节。"""
        try:
            style_name = (paragraph.style.name or "").lower()
        except Exception:  # noqa: BLE001
            return ""
        if "heading" in style_name or "标题" in style_name:
            level = "".join(ch for ch in style_name if ch.isdigit())
            return "#" * (int(level) if level else 1) + " "
        return ""

    @staticmethod
    def _render_table(table) -> list[str]:
        rows = list(table.rows)
        if not rows:
            return []
        header_cells = [cell.text.strip().replace("\n", " ") for cell in rows[0].cells]
        rendered: list[str] = ["| " + " | ".join(header_cells) + " |"]
        for row in rows[1:]:
            values = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            pairs = [
                f"{header_cells[index] or f'列{index + 1}'}: {value}"
                for index, value in enumerate(values)
                if value
            ]
            if pairs:
                rendered.append("| " + " | ".join(pairs) + " |")
        return rendered


# --------------------------------------------------------------------------- #
# 内嵌图片提取
# --------------------------------------------------------------------------- #


def _local_name(element) -> str:
    """取元素的本地标签名（忽略命名空间前缀，如 ``w:p`` -> ``p``）。"""
    tag = getattr(element, "tag", "")
    if not isinstance(tag, str):
        return ""
    return tag.rsplit("}", 1)[-1]


def _image_rel_ids(element) -> list[str]:
    """收集元素（含后代）引用的关系 ID，去重且保持出现顺序。"""
    found: list[str] = []
    for node in element.iter():
        for attribute in RELATIONSHIP_ATTRS:
            value = node.get(attribute)
            if value and value not in found:
                found.append(value)
    return found


def _build_asset(index: int, rel_id: str, related: dict, anchor: str, warnings: list[str]) -> ParsedAsset | None:
    """把关系 ID 解成图片资产；非图片关系（如超链接）返回 None。"""
    part = related.get(rel_id)
    if part is None:
        return None
    content_type = (getattr(part, "content_type", "") or "").lower()
    if not content_type.startswith("image/"):
        return None
    try:
        blob = part.blob
    except Exception as exc:  # noqa: BLE001 - 第三方部件读取异常类型不稳定
        warnings.append(f"第 {index} 张内嵌图片读取失败：{exc}")
        return None
    if not blob:
        return None
    width, height = read_size(blob)
    return ParsedAsset(
        index=index,
        data=blob,
        mime=content_type,
        extension=extension_for(content_type),
        source=ASSET_DOCX,
        placeholder=f"[图片 {index}]",
        anchor=truncate((anchor or "").strip(), 60) or "文档正文",
        width=width,
        height=height,
    )


__all__ = ["DocxParser"]
