"""Excel 需求清单解析器（V2.0 S7）。

Excel 里的「需求文档」实际有两种完全不同的形态，解析方式也必须分开：

  1. **字段约束表**（字段名 / 类型 / 必填 / 取值范围）—— 需求文档里的表，
     原样渲染成 Markdown 表格，交给规则解析器抽字段约束即可；
  2. **需求清单**（需求编号 / 需求标题 / 需求描述 / 优先级）—— 这种表**每行就是
     一条需求**。若也按表格渲染，规则解析器会把「需求编号」当成字段名，
     产出一堆没有意义的字段。因此这里逐行还原成 ``## 标题`` + ``列名: 值`` 的
     章节结构，正好对上规则解析器的章节切分。

用 openpyxl 的 read_only 模式逐行读，避免大表把内存吃满；行数与表数都设上限，
截断时如实告警。
"""
from __future__ import annotations

import io

from app.core.exceptions import DocumentParseError
from app.parsers.base import DocumentParser, ParsedDocument

try:  # pragma: no cover - 取决于运行环境
    from openpyxl import load_workbook

    EXCEL_AVAILABLE = True
except ImportError:  # pragma: no cover
    load_workbook = None  # type: ignore[assignment]
    EXCEL_AVAILABLE = False

MAX_SHEETS = 20
MAX_ROWS_PER_SHEET = 500
MAX_COLUMNS = 40

#: 命中这些列名说明这张表是「需求清单」而不是「字段约束表」。
REQUIREMENT_HINTS: tuple[str, ...] = ("需求编号", "需求标题", "需求描述", "需求名称", "需求")
TITLE_HINTS: tuple[str, ...] = ("标题", "名称", "需求")
MODULE_HINTS: tuple[str, ...] = ("模块", "所属模块", "页面")


class ExcelParser(DocumentParser):
    name = "excel"
    label = "Excel 解析器"
    extensions = ("xlsx", "xlsm")

    def parse(self, data: bytes, *, filename: str = "") -> ParsedDocument:
        if not EXCEL_AVAILABLE:  # pragma: no cover
            raise DocumentParseError(
                "解析 Excel 需要安装 openpyxl，请执行：pip install openpyxl",
                detail={"missing_dependency": "openpyxl"},
            )

        warnings: list[str] = []
        try:
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 - openpyxl 异常类型较多，统一兜底
            raise DocumentParseError(
                f"Excel 打开失败，文件可能已损坏、受密码保护，或是旧版 .xls 格式：{exc}。"
                "旧版 .xls 请先用 Excel 另存为 .xlsx 再上传。",
                detail={"filename": filename, "error": str(exc)},
            ) from exc

        lines: list[str] = []
        table_count = 0
        try:
            sheet_names = list(workbook.sheetnames)
            if len(sheet_names) > MAX_SHEETS:
                warnings.append(
                    f"工作簿含 {len(sheet_names)} 个工作表，仅解析前 {MAX_SHEETS} 个。"
                )
            for sheet_name in sheet_names[:MAX_SHEETS]:
                rows = read_rows(workbook[sheet_name], warnings, sheet_name)
                if not rows:
                    continue
                lines.append(f"# {sheet_name}")
                lines.append("")
                if looks_like_requirement_list(rows[0]):
                    lines.extend(render_requirement_rows(rows, warnings, sheet_name))
                else:
                    lines.extend(render_table(rows))
                    table_count += 1
                lines.append("")
        finally:
            workbook.close()

        if not lines:
            warnings.append("工作簿里没有读到任何非空单元格。")
        return self._finalize(
            "\n".join(lines),
            extension=self._extension(filename),
            table_count=table_count,
            metadata={"filename": filename},
            warnings=warnings,
        )

    @staticmethod
    def _extension(filename: str) -> str:
        if "." in filename:
            return filename.rsplit(".", 1)[1].lower()
        return "xlsx"


def read_rows(sheet, warnings: list[str], sheet_name: str) -> list[list[str]]:
    """读出非空行（逐行流式，避免一次性把整表加载进内存）。"""
    rows: list[list[str]] = []
    for index, raw in enumerate(sheet.iter_rows(values_only=True)):
        if index >= MAX_ROWS_PER_SHEET:
            warnings.append(
                f"工作表「{sheet_name}」超过 {MAX_ROWS_PER_SHEET} 行，已截断，超出部分未解析。"
            )
            break
        cells = ["" if value is None else str(value).strip() for value in raw[:MAX_COLUMNS]]
        if any(cells):
            rows.append(cells)
    return rows


def looks_like_requirement_list(header: list[str]) -> bool:
    """表头是否表明「每行一条需求」。"""
    joined = " ".join(header)
    return any(hint in joined for hint in REQUIREMENT_HINTS)


def render_table(rows: list[list[str]]) -> list[str]:
    """渲染成 Markdown 表格（保持字段约束表的原有解析路径）。"""
    width = max(len(row) for row in rows)
    header = [cell or f"列{index + 1}" for index, cell in enumerate(pad(rows[0], width))]
    lines = ["| " + " | ".join(header) + " |", "| " + " | ".join("---" for _ in header) + " |"]
    for row in rows[1:]:
        lines.append("| " + " | ".join(pad(row, width)) + " |")
    return lines


def render_requirement_rows(rows: list[list[str]], warnings: list[str], sheet_name: str) -> list[str]:
    """把「每行一条需求」的清单还原成章节结构。"""
    header = rows[0]
    lines: list[str] = []
    title_index = index_of(header, TITLE_HINTS)
    module_index = index_of(header, MODULE_HINTS)
    for position, row in enumerate(rows[1:], start=1):
        title = ""
        if title_index is not None and title_index < len(row):
            title = row[title_index]
        lines.append(f"## {title or f'{sheet_name} 需求 {position}'}")
        if module_index is not None and module_index < len(row) and row[module_index]:
            lines.append(f"模块：{row[module_index]}")
        for index, value in enumerate(row):
            if not value or index == title_index or index == module_index:
                continue
            lines.append(f"{header[index] or f'列{index + 1}'}：{value}")
        lines.append("")
    if len(rows) == 1:
        warnings.append(f"工作表「{sheet_name}」只有表头没有数据行。")
    return lines


def index_of(header: list[str], hints: tuple[str, ...]) -> int | None:
    for hint in hints:
        for index, cell in enumerate(header):
            if hint in (cell or ""):
                return index
    return None


def pad(row: list[str], width: int) -> list[str]:
    return [*row, *([""] * (width - len(row)))]


__all__ = [
    "ExcelParser",
    "looks_like_requirement_list",
    "read_rows",
    "render_requirement_rows",
    "render_table",
]
